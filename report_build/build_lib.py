"""Helper library: python-docx builder implementing the template's formatting rules."""
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"
HDR_FILL = "DCE6F1"
ALT_FILL = "F4F7FB"
BORDER = "BFC9D6"
CODE_FILL = "F6F6F6"
LINE_SP = 1.3

TAG_RE = re.compile(r"(</?(?:b|i|sub|sup|c)>)")


def _set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def _field(par, instr, size=12):
    def mk(t, txt=None):
        r = par.add_run()
        _set_font(r, size=size)
        if t == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        else:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), t)
        r._element.append(e)
    mk("begin")
    mk("instr", instr)
    mk("separate")
    r = par.add_run("1")
    _set_font(r, size=size)
    mk("end")


class Report:
    def __init__(self):
        self.doc = Document()
        self.toc = []        # (label, key)
        self.tables = []     # (number, title)
        self.figures = []    # (number, title)
        self.chapter_no = 0
        self.refdb = {}
        self.cited = []
        self.fig_n = 0
        self.tab_n = 0
        self.lst_n = 0
        self.eq_n = 0
        self._styles()
        self._page()

    # ------------------------------------------------------------ setup
    def _styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.line_spacing = LINE_SP
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True

    def _page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Inches(1))
        s.footer_distance = Inches(0.5)

    # ------------------------------------------------------------ sections / numbering
    def _pgnum(self, section, fmt, start):
        sect = section._sectPr
        for old in sect.findall(qn("w:pgNumType")):
            sect.remove(old)
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:fmt"), fmt)
        pg.set(qn("w:start"), str(start))
        sect.append(pg)

    def _footer(self, section, numbered):
        section.footer.is_linked_to_previous = False
        ft = section.footer
        for p in list(ft.paragraphs)[1:]:
            p._element.getparent().remove(p._element)
        p = ft.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if numbered:
            _field(p, "PAGE", 12)

    def first_section_plain(self):
        self._footer(self.doc.sections[0], False)

    def new_section(self, fmt, start=1, numbered=True):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Inches(1))
        s.footer_distance = Inches(0.5)
        self._pgnum(s, fmt, start)
        self._footer(s, numbered)
        return s

    def page_break(self):
        from docx.enum.text import WD_BREAK
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ rich text
    def _cite(self, m):
        nums = []
        for k in m.group(1).split(","):
            k = k.strip()
            if k not in self.refdb:
                raise KeyError(f"unknown reference key {k}")
            if k not in self.cited:
                self.cited.append(k)
            nums.append(self.cited.index(k) + 1)
        return ", ".join(f"[{n}]" for n in nums)

    def rich(self, par, text, size=None, base_bold=False, base_italic=False, color=None):
        text = re.sub(r"\[@([A-Za-z0-9_,\s]+)\]", self._cite, text)
        bold, ital, sub, sup, code = base_bold, base_italic, False, False, False
        for tok in TAG_RE.split(text):
            if not tok:
                continue
            m = re.fullmatch(r"<(/?)(b|i|sub|sup|c)>", tok)
            if m:
                on = m.group(1) == ""
                t = m.group(2)
                if t == "b":
                    bold = on or base_bold
                elif t == "i":
                    ital = on or base_italic
                elif t == "sub":
                    sub = on
                elif t == "sup":
                    sup = on
                elif t == "c":
                    code = on
                continue
            r = par.add_run(tok)
            if code:
                _set_font(r, MONO, (size or 12) - 1.5, bold, ital, color)
            else:
                _set_font(r, FONT, size, bold, ital, color)
            if sub:
                r.font.subscript = True
            if sup:
                r.font.superscript = True
        return par

    def p(self, text, align=None, size=None, space_after=6, indent_first=False, keep_next=False,
          bold=False, italic=False, left=None):
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_after = Pt(space_after)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left":
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif align == "right":
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if keep_next:
            pf.keep_with_next = True
        if left is not None:
            pf.left_indent = Inches(left)
        self.rich(par, text, size, bold, italic)
        return par

    def bullets(self, items, size=None, marker="•", left=0.3):
        for it in items:
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Inches(left + 0.22)
            pf.first_line_indent = Inches(-0.22)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Inches(left + 0.22))
            self.rich(par, f"{marker}\t{it}", size)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def numbered(self, items, size=None, left=0.3, fmt="({})"):
        for k, it in enumerate(items, 1):
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Inches(left + 0.3)
            pf.first_line_indent = Inches(-0.3)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Inches(left + 0.3))
            self.rich(par, f"{fmt.format(k)}\t{it}", size)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def defs(self, items, size=None, left=0.3):
        """Signal glossary: [(name, description)]."""
        for name, desc in items:
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Inches(left + 0.22)
            pf.first_line_indent = Inches(-0.22)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Inches(left + 0.22))
            self.rich(par, f"•\t<c>{name}</c> – {desc}", size)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    # ------------------------------------------------------------ headings
    def chapter(self, title, first=False):
        self.chapter_no += 1
        self.fig_n = self.tab_n = self.lst_n = self.eq_n = 0
        self.sec_marker = None
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = par.paragraph_format
        pf.page_break_before = not first
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.keep_with_next = True
        self.rich(par, f"CHAPTER {self.chapter_no}", 14, True)
        par2 = self.doc.add_paragraph()
        par2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par2.paragraph_format.space_after = Pt(16)
        par2.paragraph_format.keep_with_next = True
        self.rich(par2, title.upper(), 14, True)
        self.toc.append((f"Chapter {self.chapter_no}  {title}", f"CHAPTER {self.chapter_no}"))

    def plain_heading(self, title, key=None, toc_label=None, page_break=True):
        """Centered, bold, uppercase unnumbered heading (Acknowledgments, References ...)."""
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.page_break_before = page_break
        par.paragraph_format.space_after = Pt(16)
        par.paragraph_format.keep_with_next = True
        self.rich(par, title.upper(), 14, True)
        return par

    def h2(self, text):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = par.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(5)
        pf.keep_with_next = True
        self.rich(par, text, 12, False, True)
        return par

    def h3(self, text):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = par.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        pf.left_indent = Inches(0.0)
        self.rich(par, text, 12, True, True)
        return par

    # ------------------------------------------------------------ equations
    def eq(self, text, label=True):
        self.eq_n += 1
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_before = Pt(4)
        pf.space_after = Pt(6)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(Inches(3.135), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT)
        pf.line_spacing = 1.15
        self.rich(par, "\t", 12)
        self.rich(par, text, 12, False, True)
        if label:
            self.rich(par, f"\t({self.chapter_no}.{self.eq_n})", 12)
        return par

    # ------------------------------------------------------------ captions
    def _caption(self, text, align, keep_next, before=6, after=6):
        par = self.doc.add_paragraph()
        par.alignment = align
        pf = par.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep_next
        pf.line_spacing = 1.1
        self.rich(par, text, 11)
        return par

    # ------------------------------------------------------------ tables
    def table(self, title, headers, rows, widths, aligns=None, size=10, bold_first_col=False,
              total_row=False, caption=True, keep=True):
        num = None
        if caption:
            self.tab_n += 1
            num = f"{self.chapter_no}.{self.tab_n}"
            self.tables.append((num, title))
            self._caption(f"Table {num} {title}", WD_ALIGN_PARAGRAPH.CENTER, True, 8, 4)
        ncol = len(headers)
        t = self.doc.add_table(rows=1, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{e}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), BORDER)
            borders.append(b)
        tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for e, v in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
            m = OxmlElement(f"w:{e}")
            m.set(qn("w:w"), str(v))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tblPr.append(mar)
        for j, wv in enumerate(widths):
            t.columns[j].width = Inches(wv)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
        aligns = aligns or ["l"] * ncol

        def fill_cell(cell, text, hdr, row_idx, col_idx):
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            if hdr:
                _shade(tcPr, HDR_FILL)
            elif row_idx % 2 == 1:
                _shade(tcPr, ALT_FILL)
            par = cell.paragraphs[0]
            pf = par.paragraph_format
            pf.space_after = Pt(0)
            pf.line_spacing = 1.05
            a = aligns[col_idx]
            par.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                             "r": WD_ALIGN_PARAGRAPH.RIGHT}[a]
            if keep and row_idx < len(rows):
                pf.keep_with_next = True
            b = hdr or (bold_first_col and col_idx == 0) or (total_row and row_idx == len(rows))
            self.rich(par, str(text), size, b)

        for j, h in enumerate(headers):
            fill_cell(t.rows[0].cells[j], h, True, 0, j)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        trPr.append(th)
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        for i, row in enumerate(rows, 1):
            cells = t.add_row().cells
            for j, v in enumerate(row):
                fill_cell(cells[j], v, False, i, j)
            trPr = t.rows[-1]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.6
        return num

    # ------------------------------------------------------------ figures
    def figure(self, path, title, width_in):
        self.fig_n += 1
        num = f"{self.chapter_no}.{self.fig_n}"
        self.figures.append((num, title))
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(path, width=Inches(width_in))
        self._caption(f"Figure {num} {title}", WD_ALIGN_PARAGRAPH.CENTER, False, 2, 10)
        return num

    # ------------------------------------------------------------ code
    def code(self, title, text, size=7.5, numbered=True, caption=True, start=1):
        if caption:
            self.lst_n += 1
            self._caption(f"Listing {self.chapter_no}.{self.lst_n} {title}", WD_ALIGN_PARAGRAPH.CENTER,
                          True, 6, 3)
        lines = text.rstrip("\n").split("\n")
        w = len(str(len(lines) + start - 1))
        for k, ln in enumerate(lines, 1):
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            hang = (w + 2) * size * 0.6 / 72.0 if numbered else 0.0
            pf.left_indent = Inches(0.08 + hang)
            pf.first_line_indent = Inches(-hang)
            pf.right_indent = Inches(0.04)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = par._p.get_or_add_pPr()
            _shade(pPr, CODE_FILL)
            bdr = OxmlElement("w:pBdr")
            for side in ("left", "right"):
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "6")
                b.set(qn("w:space"), "4")
                b.set(qn("w:color"), "9AA8BA")
                bdr.append(b)
            if k == 1:
                b = OxmlElement("w:top")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "6")
                b.set(qn("w:space"), "2")
                b.set(qn("w:color"), "9AA8BA")
                bdr.append(b)
            if k == len(lines):
                b = OxmlElement("w:bottom")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "6")
                b.set(qn("w:space"), "2")
                b.set(qn("w:color"), "9AA8BA")
                bdr.append(b)
            pPr.append(bdr)
            pf.widow_control = False
            prefix = f"{k + start - 1:>{w}}  " if numbered else ""
            r = par.add_run(prefix)
            _set_font(r, MONO, size, color="7A7A7A")
            body = ln.replace("\t", "    ")
            r2 = par.add_run(body if body else " ")
            _set_font(r2, MONO, size)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.6

    # ------------------------------------------------------------ save
    def save(self, path):
        self._normalise_order()
        self.doc.save(path)

    # Word is strict about the order of child elements; python-docx appends some of ours at the end.
    _ORDER = {
        "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
                "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
                "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
                "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
                "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
                "sectPr", "pPrChange"],
        "pBdr": ["top", "left", "bottom", "right", "between", "bar"],
        "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
                 "textDirection", "tcFitText", "vAlign", "hideMark"],
        "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                  "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
                  "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
        "sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                   "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                   "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"],
    }

    def _normalise_order(self):
        roots = [self.doc.element.body]
        for sec in self.doc.sections:
            try:
                roots.append(sec.footer._element)
            except Exception:
                pass
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for root in roots:
            for tag, order in self._ORDER.items():
                rank = {W + n: i for i, n in enumerate(order)}
                for el in root.iter(W + tag):
                    kids = list(el)
                    kids_sorted = sorted(kids, key=lambda k: rank.get(k.tag, len(order)))
                    if kids != kids_sorted:
                        for k in kids:
                            el.remove(k)
                        for k in kids_sorted:
                            el.append(k)
