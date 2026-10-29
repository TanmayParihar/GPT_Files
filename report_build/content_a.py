"""Front matter, Chapter 1 and Chapter 2."""
import os
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs")

TITLE = "A Sparsity-Aware Near-Memory SpMVM Accelerator for Edge Inference on FPGA"
NAME = "Vemana Venkata Pujithram"
REG = "230907016"
GUIDE = "Dr. Rahul Ratnakumar"
DEPT = "Department of Electronics and Communication Engineering"

ABBREVS = [
    ("ADC", "Analog-to-Digital Converter"),
    ("AXI", "Advanced eXtensible Interface"),
    ("BRAM", "Block Random Access Memory"),
    ("CE", "Clock Enable"),
    ("CIM", "Compute-in-Memory"),
    ("CNN", "Convolutional Neural Network"),
    ("COO", "Coordinate (storage format)"),
    ("CSC", "Compressed Sparse Column"),
    ("CSR", "Compressed Sparse Row"),
    ("DMA", "Direct Memory Access"),
    ("DRAM", "Dynamic Random Access Memory"),
    ("DSP", "Digital Signal Processing (slice)"),
    ("ECG / EEG", "Electrocardiogram / Electroencephalogram"),
    ("FDCE", "D flip-flop with Clock Enable and asynchronous Clear"),
    ("FF", "Flip-Flop"),
    ("FPGA", "Field-Programmable Gate Array"),
    ("FSM", "Finite State Machine"),
    ("GOPS", "Giga Operations Per Second"),
    ("HBM", "High Bandwidth Memory"),
    ("IMC", "In-Memory Computing"),
    ("INT8", "8-bit signed integer"),
    ("IP", "Intellectual Property block"),
    ("LUT", "Look-Up Table"),
    ("MAC", "Multiply–Accumulate"),
    ("MLP", "Multi-Layer Perceptron"),
    ("NMC", "Near-Memory Computing"),
    ("NNZ", "Number of Non-Zero elements"),
    ("OOC", "Out-Of-Context (synthesis)"),
    ("PE", "Processing Element"),
    ("PL / PS", "Programmable Logic / Processing System"),
    ("RAMB18E1", "7-series 18 Kb Block RAM primitive"),
    ("RTL", "Register-Transfer Level"),
    ("SoC", "System on Chip"),
    ("SpMV / SpMVM", "Sparse Matrix–Vector Multiplication"),
    ("SRAM", "Static Random Access Memory"),
    ("TB", "Testbench"),
    ("XSA", "Xilinx Support Archive (hardware handoff file)"),
]


def _centered(R, text, size=12, bold=False, after=6, italic=False):
    par = R.p(text, align="center", size=size, space_after=after, bold=bold, italic=italic)
    par.paragraph_format.line_spacing = 1.15
    return par


def title_page(R):
    R.first_section_plain()
    par = R.doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(14)
    par.add_run().add_picture(os.path.join(FIG, "logo_Im1.png"), width=Inches(6.0))
    _centered(R, DEPT.upper(), 12, True, 20)
    _centered(R, "A SPARSITY-AWARE NEAR-MEMORY SpMVM", 15, True, 0)
    _centered(R, "ACCELERATOR FOR EDGE INFERENCE ON FPGA", 15, True, 18)
    _centered(R, "Minor Specialization Report", 13, True, 18)
    _centered(R, "Submitted by", 12, False, 8, True)
    _centered(R, NAME, 13, True, 0)
    _centered(R, f"Registration No.: {REG}", 12, False, 4)
    _centered(R, "(Individual Minor Specialization Project)", 11, False, 18, True)
    _centered(R, "Under the guidance of", 12, False, 4, True)
    _centered(R, GUIDE, 13, True, 0)
    _centered(R, f"{DEPT}, MIT Manipal", 11.5, False, 18)
    _centered(R, "in partial fulfilment of the requirements for the award of the degree of", 12, False, 4)
    _centered(R, "MINOR SPECIALIZATION IN", 12.5, True, 2)
    _centered(R, "Embedded Systems", 13, True, 16)
    par = R.doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(6)
    par.add_run().add_picture(os.path.join(FIG, "logo_Im2.png"), width=Inches(0.55))
    _centered(R, "MANIPAL INSTITUTE OF TECHNOLOGY", 13, True, 0)
    _centered(R, "(A Constituent Unit of Manipal Academy of Higher Education)", 11, False, 0)
    _centered(R, "MANIPAL-576104, KARNATAKA, INDIA", 11, False, 4)
    _centered(R, "October 2026", 11.5, False, 0)


def acknowledgments(R):
    R.new_section("lowerRoman", 1, True)
    R.plain_heading("Acknowledgments", page_break=False)
    R.p("I express my sincere gratitude to my project guide, Dr. Rahul Ratnakumar, Department of Electronics "
        "and Communication Engineering, Manipal Institute of Technology, Manipal, for his constant guidance, "
        "critical feedback and encouragement throughout this work. His insistence on connecting every "
        "architectural decision to a measurable hardware quantity shaped the way this project was formulated, "
        "implemented and reported, and the discussions that led to the conference paper on which this report "
        "builds were invaluable.")
    R.p("I thank the Head of the Department, the project coordinators and the faculty of the Department of "
        "Electronics and Communication Engineering, MIT Manipal, for the academic environment, the laboratory "
        "facilities and the software licences that made the SystemVerilog design, Vivado simulation and "
        "synthesis experiments possible. I also thank the laboratory staff for their technical assistance with "
        "the development environment.")
    R.p("I am grateful to my classmates and friends for the many discussions on sparse linear algebra, FPGA "
        "memory resources and verification practice, and to my family for their patience and support during "
        "the semester. Finally, I acknowledge the authors of the open literature and the vendor documentation "
        "cited in this report, on whose work the background and the comparisons presented here depend.")
    R.p(f"<b>{NAME}</b>", align="right")


def abstract(R):
    R.plain_heading("Abstract")
    R.p("Deploying deep neural networks on edge devices is limited by the cost of memory access and by the "
        "wasted arithmetic spent on the zero weights that survive pruning. This project presents the design, "
        "verification and post-synthesis characterisation of a Sparsity-Aware Near-Memory Sparse "
        "Matrix–Vector Multiplication (SpMVM) accelerator for edge inference on an AMD/Xilinx Zynq-7000 "
        "SoC (xc7z020clg484-1). True analog in-memory computing relies on access to SRAM bit-lines that "
        "commercial FPGA fabrics do not expose; the architecture therefore realises the practical FPGA "
        "equivalent, BRAM-centric near-memory computing, by placing a zero-skipping controller and a signed "
        "INT8 multiply–accumulate (MAC) unit directly at the synchronous read ports of three on-chip block "
        "RAMs that hold the compressed weights, the column indices and the dense input vector.")
    R.p("Weights are stored in Compressed Sparse Row (CSR) form so that only non-zero elements are fetched and "
        "multiplied. A five-state finite state machine (IDLE, FETCH_PTR, FETCH_VEC, COMPUTE, DONE_STATE) "
        "sequences each stored non-zero in exactly three clock cycles, absorbing the one-cycle BRAM read "
        "latency of the column-index bank that feeds the indirect address of the vector bank; the resulting "
        "latency is T = 1 + 3·NNZ + 1 cycles. The arithmetic unit multiplies two signed 8-bit operands "
        "into an exact 16-bit product and accumulates into a 32-bit register without DSP48E1 slices. The "
        "design is verified with a directed MAC testbench and a top-level testbench that preloads the BRAMs and "
        "reproduces the expected accumulation 0 → 15 → 7 for the dot product 5(3) + (−4)(2).")
    R.p("Out-of-context synthesis on the Zynq-7000 device reports 154 slice LUTs, 41 slice registers (FDCE), "
        "3 RAMB18E1 block RAMs and 0 DSP48E1 slices, a 10.32 ns critical path delay (about 96.9 MHz), and a "
        "vectorless post-synthesis power estimate of 4.502 W total on-chip power (4.235 W dynamic, 0.268 W "
        "static). At 90 % sparsity the core delivers an effective dense-equivalent throughput of about "
        "0.64 GOPS. The report also derives storage and latency models for Dense, CSR, COO and CSC formats and "
        "lays out a system-level roadmap covering multi-row CSR sequencing, AXI4-Lite wrapping, Zynq PS–PL "
        "integration and board deployment.")
    R.p("<b>Keywords:</b> FPGA, Sparse Matrix–Vector Multiplication, CSR, Near-Memory Computing, "
        "Compute-in-Memory, Block RAM, INT8 MAC, Zynq-7000, SystemVerilog.", align="left")


def _list_table(R, headers, rows, widths):
    R.table("", headers, rows, widths, aligns=["l", "l", "l"][:len(headers)], size=11, caption=False)


def contents(R, pages):
    R.plain_heading("Table of Contents")
    items = [
        ("Acknowledgments", "Acknowledgments"), ("Abstract", "Abstract"),
        ("List of tables", "List of tables"), ("List of figures", "List of figures"),
        ("List of abbreviations", "List of abbreviations"),
        ("Chapter 1  Introduction", "CHAPTER 1"),
        ("Chapter 2  Literature review and objectives", "CHAPTER 2"),
        ("Chapter 3  Methodology", "CHAPTER 3"),
        ("Chapter 4  Contribution of student", "CHAPTER 4"),
        ("Chapter 5  Result analysis", "CHAPTER 5"),
        ("Chapter 6  Conclusion and future scope", "CHAPTER 6"),
        ("Visible outcome", "Visible outcome"), ("References", "References"),
        ("Minor specialization project details", "Minor specialization project details"),
    ]
    rows = [[a, str(pages.get(k, "–"))] for a, k in items]
    R.table("", ["Section", "Page"], rows, [5.4, 0.87], aligns=["l", "l"], size=11, caption=False, keep=False)


def list_of_tables(R, reg, pages):
    R.plain_heading("List of Tables")
    rows = [[n, t, str(pages.get("T:" + n, "–"))] for n, t in reg["tables"]] or [["–", "–", "–"]]
    R.table("", ["No", "Table title", "Page"], rows, [0.55, 4.97, 0.75], size=10.5, caption=False, keep=False)


def list_of_figures(R, reg, pages):
    R.plain_heading("List of Figures")
    rows = [[n, t, str(pages.get("F:" + n, "–"))] for n, t in reg["figures"]] or [["–", "–", "–"]]
    R.table("", ["No", "Figure title", "Page"], rows, [0.55, 4.97, 0.75], size=10.5, caption=False, keep=False)


def abbreviations(R):
    R.plain_heading("List of Abbreviations")
    R.table("", ["Abbreviation", "Expansion"], [[a, b] for a, b in ABBREVS], [1.5, 4.77], size=10.5,
            caption=False, keep=False)


# =====================================================================================
def chapter1(R):
    R.new_section("decimal", 1, True)
    R.chapter("Introduction", first=True)

    R.h2("1.1 Background and problem statement")
    R.p("Matrix–vector multiplication is the central computation of neural-network inference: every "
        "fully connected layer, the projection layers of a Transformer, and the lowered form of a convolution "
        "all reduce to evaluating <i>y</i> = <i>A</i><i>x</i>, where <i>A</i> is an <i>M</i> × <i>N</i> "
        "weight matrix and <i>x</i> is an <i>N</i>-element activation vector [@jouppi2017]. The elements of "
        "the result are <i>y</i><sub><i>i</i></sub> = ∑<sub><i>j</i></sub> <i>A</i><sub><i>ij</i></sub> "
        "<i>x</i><sub><i>j</i></sub>. After training, weight pruning removes connections whose magnitude is "
        "small, and sparsity levels of 50–90 % or higher are routinely reached with limited accuracy loss "
        "[@han2015], [@hoefler2021]. The pruned matrix <i>A</i> therefore contains mostly zeros.")
    R.p("A conventional dense datapath ignores this structure. It stores every zero, fetches every zero from "
        "memory and spends a clock cycle and the associated dynamic energy multiplying by it [@sze2017]. The "
        "waste has two components. The first is arithmetic: for a matrix of density ρ = NNZ/(<i>MN</i>), "
        "a fraction 1 − ρ of all multiply–accumulate operations has no effect on the result. The "
        "second, and usually more expensive, is data movement. Fetching one weight from off-chip DRAM costs "
        "orders of magnitude more energy than the multiplication that consumes it [@horowitz2014], [@sze2017]; "
        "on an edge device with a battery or a few-watt thermal budget this “memory wall”, rather than "
        "the arithmetic itself, bounds the achievable inferences per joule.")
    R.p("The engineering problem addressed in this project is therefore stated as follows: design a compact "
        "hardware engine for <i>y</i> = <i>A</i><i>x</i> with a pruned, quantised <i>A</i> that (i) keeps all "
        "operand traffic on-chip, (ii) performs work proportional to the number of non-zero weights rather than "
        "to <i>MN</i>, and (iii) maps onto the resources of a commercial low-cost FPGA with a minimal "
        "arithmetic footprint. The target platform is the AMD/Xilinx Zynq-7000 SoC "
        "(xc7z020clg484-1) [@ug585], [@ds190], which offers 53,200 LUTs, 106,400 flip-flops, 280 RAMB18E1 "
        "block RAMs and 220 DSP48E1 slices alongside a dual-core ARM Cortex-A9 processing system.")

    R.h2("1.2 Motivation and relevance")
    R.p("Four properties of FPGA-based inference motivate the approach taken here. <b>Determinism:</b> "
        "an FPGA executes a fixed finite state machine against synchronous memories, so the latency of an "
        "inference is a closed-form function of the stored data and not of cache state, operating-system "
        "scheduling or DRAM refresh. This is the property that makes FPGAs attractive for in-vehicle, "
        "robotic and medical workloads with hard real-time deadlines. <b>On-chip locality:</b> the Zynq-7020 "
        "provides about 4.9 Mb of block RAM distributed in columns across the fabric [@ug473]; a pruned "
        "layer tile held in these RAMs never touches DRAM during inference. <b>Low precision:</b> INT8 "
        "quantisation of weights and activations is well established for inference "
        "[@jacob2018], [@gholami2021], and an 8 × 8 signed multiplier is small enough to be built from "
        "look-up tables, avoiding the DSP48E1 slices that are sized for 25 × 18-bit products "
        "[@ug479]. <b>Sparsity:</b> storing and processing only non-zero weights converts the pruning ratio "
        "directly into reduced storage, reduced memory traffic and reduced cycle count.")
    R.p("A further motivation is conceptual. The research literature on <i>in-memory computing</i> (IMC) "
        "[@verma2019], [@sebastian2020], [@jhang2021] describes computing inside the memory array itself, "
        "which removes the weight-movement cost entirely. True analog IMC is, however, a property of custom "
        "SRAM or non-volatile memory macros whose bit-lines are accessible to the designer. On a commercial "
        "FPGA the block RAM is a hard primitive exposing only a synchronous address–data interface. The "
        "closest practical realisation is <i>near-memory computing</i> (NMC): placing the compute logic "
        "immediately at the memory read port so that every operand is consumed in the cycle after it is "
        "read, with no intermediate buffering and no off-chip transfer. Section 3.5 develops this distinction "
        "rigorously, and the design in this report is a BRAM-centric NMC engine in that sense.")

    R.h2("1.3 Main engineering contribution")
    R.p("The contribution of this project is a modular, parameterisable <b>BRAM-centric Sparse "
        "Compute-in-Memory / Near-Memory Processing Engine</b> (the Sparse-CIM PE, top module "
        "<c>sparse_cim_top</c>) together with its complete verification and hardware characterisation. The "
        "engine integrates the following elements:")
    R.bullets([
        "<b>Compressed weight and index banks.</b> Two <c>bram_bank</c> instances (<c>bank_values</c> and "
        "<c>bank_col_idx</c>) hold the non-zero INT8 weights and their column indices, addressed by a common "
        "non-zero pointer.",
        "<b>Indirect vector fetching.</b> A third bank, <c>bank_vector</c>, holds the dense input vector and is "
        "read at the address delivered by the column-index bank, so that the activation matching each "
        "non-zero weight is fetched without scanning the vector.",
        "<b>A five-state zero-skipping controller.</b> <c>sparse_cim_controller</c> sequences IDLE, FETCH_PTR, "
        "FETCH_VEC, COMPUTE and DONE_STATE; because it iterates over <c>nnz_count</c> stored non-zeros, zero "
        "weights consume neither cycles nor memory reads.",
        "<b>A DSP-free signed INT8 MAC.</b> <c>mac_unit</c> forms an exact 16-bit signed product and "
        "accumulates into a 32-bit register under an enable and a synchronous clear, and synthesises entirely "
        "into LUTs and flip-flops.",
        "<b>Verification and characterisation.</b> A directed MAC testbench and a top-level testbench "
        "demonstrate correct signed accumulation, and Vivado out-of-context synthesis quantifies area "
        "(154 LUTs, 41 FFs, 3 RAMB18E1, 0 DSPs), delay (10.32 ns) and vectorless power (4.502 W).",
    ])

    R.h2("1.4 Scope and evolution of the design")
    R.p("The project is organised around a mathematical formulation, an implemented compute core and a "
        "scaling framework. The mathematical formulation (Section 3.4) defines the CSR representation "
        "of <i>A</i> by its <c>values</c>, <c>col_idx</c> and <c>row_ptr</c> arrays and proves the equivalence "
        "of the zero-skipping computation to the dense product. The implemented core evaluates the dot "
        "product of one compressed row with the dense vector and accumulates it in the 32-bit register; "
        "this is the unit that is simulated, synthesised and characterised. The scaling framework "
        "(Sections 3.10, 3.13 and 6.4) extends the same datapath with <c>row_ptr</c> sequencing, result "
        "write-back, an AXI4-Lite register interface and Zynq PS–PL integration. Table 1.1 summarises how "
        "these layers relate.")
    R.table("Layers of the design and their role in this report",
            ["Layer", "Architecture described", "Role in this report"],
            [["Mathematical model", "<i>y</i> = <i>A</i><i>x</i> with CSR arrays <c>values</c>, <c>col_idx</c>, "
              "<c>row_ptr</c>; zero-skipping dot products", "Formulation, storage and cycle models (Sections 3.4, 3.6, 5.9)"],
             ["Implemented core", "Three BRAM banks, 5-state FSM, INT8 MAC with 32-bit accumulator; one "
              "compressed row per invocation", "Simulated, synthesised and characterised (Chapters 3 and 5)"],
             ["Scaling framework", "6-state multi-row CSR FSM: IDLE → FETCH_ROW_PTR → STREAM_CSR "
              "→ MAC_ACCUM → WRITEBACK → DONE", "Architectural extension (Sections 3.10, 3.13)"],
             ["SoC integration", "AXI4-Lite wrapper, AXI BRAM controllers, Zynq PS (ARM Cortex-A9), "
              ".xsa export, Vitis driver", "Integration roadmap (Sections 3.20–3.23, 6.4)"]],
            [1.3, 3.0, 1.97], size=10)

    R.h2("1.5 Demonstration boundary")
    R.p("The demonstrated results of this project are of two kinds. <b>Behavioural simulation</b> in the "
        "Vivado Simulator exercises the RTL with directed stimulus and shows the expected accumulation and the "
        "<c>done</c> handshake. <b>Post-synthesis characterisation</b> on the Zynq-7000 device "
        "(<c>xc7z020clg484-1</c>) reports resource utilisation, critical-path delay and vectorless power for "
        "the out-of-context core, which isolates the accelerator from I/O pads and from the rest of a "
        "system. Together these establish the correctness of the datapath and the hardware cost of the "
        "engine.")
    R.p("System-level deployment — the AXI4-Lite register interface, the Zynq processing-system "
        "integration, bitstream export and execution on a physical board — forms the integration "
        "roadmap of the project. It is specified in Sections 3.20–3.23 so that the next phase can proceed "
        "directly from this report, and it is the natural continuation of the work rather than part of the "
        "verified results presented in Chapter 5.")


# =====================================================================================
def chapter2(R):
    R.chapter("Literature Review and Objectives")

    R.h2("2.1 Pruned neural networks and sparse inference")
    R.p("Han <i>et al.</i> showed that a trained network can be pruned by removing low-magnitude weights and "
        "retraining, reducing the parameter count of large networks by an order of magnitude without loss of "
        "accuracy [@han2015]. Subsequent work established a taxonomy of sparsity patterns [@hoefler2021]. "
        "<i>Unstructured</i> sparsity removes individual weights anywhere in a matrix; it achieves the "
        "highest compression for a given accuracy but yields irregular access patterns that dense hardware "
        "cannot exploit. <i>Structured</i> sparsity removes whole rows, columns, filters or fixed-size "
        "blocks, or enforces a pattern such as 2:4 (two non-zeros in every group of four), which regularises "
        "the index stream at some cost in flexibility [@mishra2021].")
    R.p("The three principal layer types differ in how sparsity appears. In <b>fully connected (FC)</b> and "
        "multi-layer perceptron layers the weight matrix is used directly as <i>A</i>, and unstructured "
        "pruning ratios above 90 % have been reported for large FC layers [@han2015], so SpMVM is the "
        "exact computation. In <b>convolutional (CNN)</b> layers, filters are lowered to matrix form "
        "(im2col) or processed with a sparse dataflow such as the Cartesian-product scheme of SCNN "
        "[@parashar2017]; both reduce to products of sparse weight vectors with dense activation windows. "
        "In <b>recurrent and Transformer</b> layers, the gate and projection matrices are large dense-"
        "activation products to which magnitude pruning applies in the same way, and the resulting "
        "matrix–vector products dominate single-batch edge inference, where the activation vector is "
        "dense and the weight matrix is sparse — precisely the SpMVM case treated in this report.")

    R.h2("2.2 Efficient sparse accelerators and memory-centric computing")
    R.p("<b>EIE.</b> The Efficient Inference Engine of Han <i>et al.</i> is an ASIC accelerator that operates "
        "directly on a compressed network [@han2016]. It stores the pruned, weight-shared matrix in an "
        "interleaved compressed-column format in on-chip SRAM distributed across processing elements, and "
        "skips zero activations dynamically. Its principal lesson for the present work is that keeping the "
        "entire compressed model in on-chip memory removes DRAM from the inference path, and that the "
        "index and pointer overhead of the compressed format must be accounted for in the energy model.")
    R.p("<b>SCNN.</b> Parashar <i>et al.</i> describe an accelerator for compressed-sparse CNNs that stores "
        "both weights and activations in compressed form and multiplies all non-zero weight–activation "
        "pairs in a Cartesian-product array, reporting performance and energy gains over a comparably "
        "provisioned dense accelerator [@parashar2017]. The lesson is that sparsity must be exploited "
        "<i>and</i> the resulting irregular accumulation must be managed in hardware.")
    R.p("<b>GraphLily and SERPENS.</b> GraphLily is an overlay for graph linear algebra on HBM-equipped FPGAs "
        "that implements SpMV and related kernels using 19 HBM channels [@hu2021]. SERPENS is an HBM-based "
        "SpMV accelerator using 24 HBM channels and wide streaming processing elements to reach high "
        "throughput [@song2022]. Both target data-centre cards: in the comparison of this report they are "
        "listed at 43 W and 48 W respectively. Their lesson is that SpMV throughput on large matrices is "
        "bandwidth-limited and that HBM parallelism is the data-centre answer — an answer unavailable "
        "within an edge power budget.")
    R.p("<b>Cerberus.</b> Hwang <i>et al.</i> propose a triple-mode accelerator that switches between "
        "compressed, bitmap and dense representations according to the sparsity of the operand, trading "
        "format overhead against skipping efficiency [@hwang2024]; it is listed at 17.032 W and 112.4 GOPS "
        "in the comparison. The lesson is that no single format is optimal across sparsity levels, which "
        "motivates the break-even analysis of Section 3.4.")
    R.p("<b>Zhang <i>et al.</i> and Abhiram <i>et al.</i></b> Zhang <i>et al.</i> present a 109-GOPS/W FPGA "
        "Vision Transformer accelerator with a weight-loop dataflow, reporting about 348 GOPS at 3.2 W "
        "using roughly 120K LUTs [@zhang2024]; it is the closest INT8 FPGA full-system comparison. "
        "Abhiram <i>et al.</i> report a fault-resilient INT8 MAC unit for mission-critical AI hardware at "
        "2967 µm<sup>2</sup>, 0.507 mW and 3.99 ns [@abhiram2024]; it is an ASIC arithmetic-only "
        "comparison point for the MAC datapath. Survey literature on FPGA sparse accelerators "
        "[@liu2024], [@leapspace2026], on MAC design [@kumar2014], [@kumar2023], and on SpMV optimisation "
        "for general-purpose processors [@williams2007], [@bell2009] frames these designs.")
    R.p("<b>Dataflow and FPGA streaming designs.</b> Eyeriss introduces a row-stationary dataflow that "
        "minimises the movement of weights, activations and partial sums by keeping reusable data in local "
        "storage next to the arithmetic [@chen2016]; FINN builds streaming FPGA inference pipelines in which "
        "each layer has its own datapath and on-chip memory [@umuroglu2017]. Neither is a sparse SpMV "
        "engine, but both support the design principle adopted here: place the memory that holds the "
        "operands directly beside the arithmetic unit that consumes them.")
    R.h3("ASIC in-memory computing versus FPGA BRAM near-memory computing")
    R.p("In ASIC <b>analog SRAM IMC</b>, multiple word-lines are activated simultaneously "
        "and the bit-line current or charge sums the products of stored weights and applied inputs; an ADC at "
        "the column periphery digitises the result [@verma2019], [@jhang2021]. Compute and storage are the "
        "same physical structure, so weights do not move at all. The price is a custom bit-cell and "
        "periphery (word-line drivers, ADCs, calibration), limited precision and sensitivity to process "
        "variation. In <b>FPGA BRAM near-memory computing</b>, the memory is a fixed synchronous macro "
        "whose bit-lines are inaccessible; computation happens in programmable logic that is adjacent to "
        "the read port. Weights still move, but only a few hundred picoseconds to a nanosecond of "
        "programmable routing, and only on-chip. The result is exact digital arithmetic, deterministic "
        "timing and full reconfigurability at the price of a read-out energy that analog IMC avoids.")

    R.h2("2.3 Numerical precision for edge inference")
    R.p("Integer-only inference with 8-bit weights and activations preserves accuracy for most vision and "
        "language models when quantisation-aware calibration is used [@jacob2018], [@gholami2021]. In a signed "
        "two's-complement INT8 representation each operand lies in [−128, 127]. The product of two "
        "INT8 operands lies in [−16,256, 16,384]: the extreme values are (−128)(127) = −16,256 "
        "and (−128)(−128) = +16,384. Both fit in a 16-bit signed number (range −32,768 to "
        "32,767), so a 16-bit product is <i>exact</i> and no rounding or truncation occurs at the multiplier "
        "output.")
    R.p("The accumulator must hold a sum of up to NNZ such products. Since |product| ≤ 2<sup>14</sup>, a "
        "32-bit signed accumulator (|value| ≤ 2<sup>31</sup> − 1) cannot overflow for fewer than "
        "2<sup>17</sup> = 131,072 terms; the addressing of the present instance (6-bit pointers, at most 63 "
        "non-zeros per invocation) leaves a margin of more than 2,000×. The trade-off against narrower "
        "accumulators is a handful of extra flip-flops and carry-chain slices, which is negligible against "
        "the guarantee of overflow-free accumulation for any practical layer tile.")

    R.h2("2.4 FPGA platform and implementation literature")
    R.p("The Zynq-7000 family integrates a dual-core ARM Cortex-A9 processing system (PS) with 7-series "
        "programmable logic (PL) on one die [@ug585]. The xc7z020 device provides 53,200 six-input LUTs, "
        "106,400 flip-flops, 140 36 Kb block RAMs (equivalently 280 RAMB18E1) and 220 DSP48E1 slices "
        "[@ds190]. Each <b>RAMB18E1</b> is an 18 Kb true dual-port synchronous memory whose output has a "
        "one-cycle read latency (optionally two with the output register enabled) and that can be "
        "configured from 16K × 1 up to 512 × 36 and, for byte-oriented data, as 2K × 9 [@ug473]. "
        "Because the read is registered inside the macro, the BRAM presents data one clock edge after "
        "the address — the fact on which the three-cycle schedule of Section 3.11 rests.")
    R.p("A <b>DSP48E1</b> slice contains a 25 × 18-bit signed multiplier, a 48-bit accumulator and a "
        "pre-adder [@ug479]. For 8-bit operands most of this capability is unused; an 8 × 8 signed "
        "multiplier built from LUTs and carry chains [@ug474] costs on the order of one hundred LUTs, "
        "leaves the DSP columns free and, because LUTs are placed anywhere in the fabric, allows the "
        "multiplier to sit adjacent to a BRAM column without the placement constraints of the DSP columns. "
        "Vivado synthesis infers block RAM from the <c>ram_style</c> attribute and from the synchronous-read "
        "coding style, and infers LUT/carry arithmetic when DSP use is not requested [@ug901]. Timing is "
        "reported against constraints [@ug903] and power is estimated either vectorlessly or from switching "
        "activity [@ug907]; Section 3.17–3.19 and Chapter 5 apply these flows.")

    R.h2("2.5 Project objectives")
    R.p("The literature review leads to the following objectives for the project:")
    R.numbered([
        "Formulate SpMVM <i>y</i> = <i>A</i><i>x</i> over CSR data and derive storage, traffic and cycle models "
        "for Dense, CSR, COO and CSC formats.",
        "Design a BRAM-centric near-memory Sparse-CIM PE in SystemVerilog comprising three parameterised BRAM "
        "banks, a five-state zero-skipping FSM and a signed INT8 MAC with 32-bit accumulation.",
        "Derive the exact cycle-accurate execution latency of the controller, including the one-cycle "
        "synchronous BRAM read latency.",
        "Verify the arithmetic and the control flow by directed simulation, including a negative weight and "
        "the <c>start</c>/<c>done</c> handshake.",
        "Characterise area, critical-path delay and vectorless power of the core on the Zynq-7000 "
        "<c>xc7z020clg484-1</c> using Vivado out-of-context synthesis, with no DSP48E1 usage.",
        "Compare the proposed engine with published sparse accelerators and MAC units on a common table of "
        "scope, precision, area, power, delay and throughput.",
        "Specify a system-level integration roadmap: multi-row CSR sequencing, AXI4-Lite IP packaging, Zynq "
        "PS–PL integration, Vitis control software and edge deployment scenarios.",
    ], fmt="{}.")
    R.table("Literature lessons and corresponding project objectives",
            ["Source", "Lesson drawn", "Project objective"],
            [["Han <i>et al.</i> (pruning) [@han2015]", "Pruned layers are mostly zeros; sparsity is exploitable",
              "Objective 1 (CSR formulation)"],
             ["EIE [@han2016]", "Keep the compressed model in on-chip memory; count index overhead",
              "Objectives 1, 2 (BRAM banks, storage model)"],
             ["SCNN [@parashar2017]", "Manage irregular accumulation of sparse products in hardware",
              "Objectives 2, 4 (FSM and MAC)"],
             ["GraphLily, SERPENS [@hu2021], [@song2022]", "HBM bandwidth is a data-centre answer; edge needs "
              "on-chip locality", "Objectives 2, 5, 6 (on-chip NMC, power)"],
             ["Cerberus [@hwang2024]", "Format choice depends on sparsity level",
              "Objective 1 (break-even analysis)"],
             ["Zhang <i>et al.</i> [@zhang2024]", "INT8 FPGA systems reach high throughput at large area",
              "Objective 6 (comparison)"],
             ["Abhiram <i>et al.</i> [@abhiram2024]", "INT8 MAC reference at ASIC level", "Objectives 5, 6"],
             ["Quantisation [@jacob2018], [@gholami2021]", "INT8 preserves accuracy; wide accumulation needed",
              "Objectives 2, 4 (16-bit product, 32-bit sum)"],
             ["Xilinx guides UG473, UG901, UG907 [@ug473], [@ug901], [@ug907]",
              "RAMB18E1 timing, synthesis inference, power flow", "Objectives 3, 5"],
             ["Zynq TRM, AXI [@ug585], [@amba]", "PS–PL integration via AXI4-Lite and BRAM controllers",
              "Objective 7 (roadmap)"]],
            [1.9, 2.5, 1.87], size=9.5)
